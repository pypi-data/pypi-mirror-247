import configparser
import re
import csv
import os
import json
import codecs
import numpy as np
import datetime
from tqdm import tqdm
import collections
import xlrd
import xlsxwriter
import docx
from docx.document import Document
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import _Cell, Table
from docx.text.paragraph import Paragraph
from win32com import client as wc
import warnings

def doc2docx(file_path, out_path=None):
    """
    doc文档转换为docx文档，默认以原路径保存，也可以自定义保存路径
    如果是doc结尾的，则转换，否则不转换
    返回转换后的路径，如果未转换，则返回原路径，路径都是获取的绝对路径
    :param file_path: 文件路径
    :param out_path: 转换为docx后的文件保存路径，默认保存到原始路径
    :return: 转换后的路径
    """
    # 获取绝对路径，因为这里用相对路径会报错
    file_path = os.path.abspath(file_path)
    file_dir = os.path.dirname(file_path)
    files = os.listdir(file_dir)
    if out_path is not None:
        if out_path.endswith('.docx'):
            save_path = os.path.abspath(out_path)
        else:
            warnings.warn('保存路径：{} 必须是docx文件名称！'.format(out_path))
            return out_path
    else:
        save_path = file_path
    if file_path.lower().endswith('.doc'):
        if out_path is None:
            save_path = file_path[:-3] + 'docx'
        save_name = os.path.basename(save_path)
        if save_name in files:
            warnings.warn('文件：{} 已存在！'.format(save_path))
            return save_path
        word = wc.Dispatch('word.Application')
        doc = word.Documents.Open(file_path)
        doc.SaveAs(save_path, 12, False, '', True, '', False, False, False, False)
        doc.Close()
        word.Quit()
    return save_path

def read_config(conf_path):
    """
    根据配置文件的路径读取配置文件，并返回配置文件内容组成的dict
        配置文件格式(名称如config.conf)：
            [conn_config]
            # sql连接配置
            host=172.19.50.66
            port=5432
            user=fpcdpc
            password=PASSWORD
            database=dpc_db
    :param conf_path: 配置文件路径
    :return: 配置文件组成的dict
    """
    conf_dict = {}
    cf = configparser.ConfigParser()
    cf.read(conf_path, encoding='utf-8')
    secs = cf.sections()
    for s in secs:
        items = cf.items(s)
        for i in items:
            conf_dict[i[0]] = i[1]
    return conf_dict


def check_and_creat_dir(dir_path):
    '''
    判断目录是否存在，目录不存在则创建目录
    :param dir_path: 目录路径
    :return:目录不存在则返回False， 存在True
    '''
    if not os.path.exists(dir_path):
        os.makedirs(dir_path)
        return False
    else:
        return True


def get_PolygonArea(points):
    """
    计算多边形面积
    :param points: [[x1, y1], [x2, y2], [x3, y3], [x4, y4], ...]
    :return: 面积
    """

    sizep = len(points)
    if sizep<3:
        return 0.0

    area = points[-1][0] * points[0][1] - points[0][0] * points[-1][1]
    for i in range(1, sizep):
        v = i - 1
        area += (points[v][0] * points[i][1])
        area -= (points[i][0] * points[v][1])

    return abs(0.5 * area)

def get_bracketed_content(text):
    """
    获取文本中所有小括号中的内容组成的list
    如：
        香港特(别行）政区)（北京点）
        return：
            ['别行', '北京点']
    :param text: 文本
    :return: 括号中内容组成的list
    """
    res = re.findall(r'[（(](.*?)[）)]', text)
    return res


def extract_base(pattern, text):
    """ 正则抽取器的基础函数
    Args:
        pattern(re.compile): 正则表达式对象
        text(str): 字符串文本

    Returns:
        list: 返回结果
    """
    results = [{'text': item.group(1),
                'offset': (item.span()[0] - 1, item.span()[1] - 1)}
               for item in pattern.finditer(text)]

    return results

def get_email(text):
    """ 提取文本中的 E-mail，返回结果为list
        如：
            text = '我的email是ifee@baidu.com和dsdsd@dsdsd.com,李林的邮箱是eewewe@gmail.com'
            emails = get_email(text)
            print(emails)
            [{'text': 'ifee@baidu.com', 'offset': (8, 22), 'domain_name': 'baidu'},
             {'text': 'dsdsd@dsdsd.com', 'offset': (23, 38), 'domain_name': 'dsdsd'},
             {'text': 'eewewe@gmail.com', 'offset': (45, 61), 'domain_name': 'gmail'}]
        其中：
            text：提取的email
            offset：email在文本中的位置，含头不含尾
            domain_name：email的域名，即@后面的第一部分
    """
    #原
    # 邮箱中，允许中文等字符存在。但是，中文字符、#、* 会对结果造成较大干扰故忽略
    # EMAIL_PATTERN = r'(?<=[^0-9a-zA-Z.\-])' \
    #                 r'([a-zA-Z0-9_.-]+@[a-zA-Z0-9_.-]+(?:\.[a-zA-Z0-9_.-]+)*\.[a-zA-Z0-9]{2,6})' \
    #                 r'(?=[^0-9a-zA-Z.\-])'
    # #改
    EMAIL_PATTERN = r'(?<=[^0-9a-zA-Z.\-])' \
                    r'([a-zA-Z0-9_.-]+@[a-zA-Z]+[a-zA-Z0-9.\-]*(?:\.[a-zA-Z0-9\-]+)*[a-zA-Z0-9]+\.[a-zA-Z0-9]{2,6})' \
                    r'(?=[^0-9a-zA-Z.\-])'
    # 抽取邮箱的域名
    # EMAIL_DOMAIN_PATTERN = r'(?<=@)([0-9a-zA-Z]+)(?=\.)'

    EMAIL_DOMAIN_PATTERN = r'(?<=@)([0-9a-zA-Z\-]+)(?=\.)'
    email_pattern = re.compile(EMAIL_PATTERN)
    text = text.replace(' at ', '@').replace(' dot ', '.')
    text = ''.join(['#', text, '#'])
    results = extract_base(email_pattern, text)
    email_domain_pattern = re.compile(EMAIL_DOMAIN_PATTERN)
    detail_results = []
    for item in results:
        domain_name = email_domain_pattern.search(item['text']).group(1)
        item.update({'domain_name': domain_name})
        detail_results.append(item)
    return detail_results


def rm_bracketed(text):
    """
    去除文本中的括号，包括括号中的内容，返回去括号后的文本
    如：
        香港特(别行）政区（北京点）
        return：
            香港特政区
    :param text:文本
    :return:去括号后的文本
    """
    res = re.sub(u"[（(](.*?)[）)]|\{.*?\}|\[.*?\]|\<.*?\>", "", text)
    return res


def rm_symbol(text):
    """
    去除文本中的所有符号，返回去符号后的文本
    如：
        香港特(别·行）政，区（北京-点）
        return：
            香港特别行政区北京点
    :param text:
    :return:
    """
    res = re.sub(
        "[\s+\.\!\/_, $%^*(+\"\')]|[ \t\r\n\\\\+—－\-()?【】“”！，。？:：、~@#￥%……&*（）\|「」▏·`▪•۰・●⁺°～’\[\➕;〔〕《–‖﹢〖〗‘》［］◆❤×『\]』｡×\\\️=；²∙﹙′★◎〉─③ⅳ―☆㎡〇ⅲ⊂♡⑧℃⑤︱╮₂ⅴⅱ³»①〈╭✘ ※❥･﹚､ⅰ<>›ܶ│丨‧丶]",
        "", text)
    return res

def LCS(x, y):
    c = np.zeros((len(x)+1, len(y)+1))
    b = np.zeros((len(x)+1, len(y)+1))
    for i in range(1, len(x)+1):
        for j in range(1, len(y)+1):
            if x[i-1] == y[j-1]:
                c[i, j] = c[i-1, j-1]+1
                b[i, j] = 2
            else:
                if c[i-1, j] >= c[i, j-1]:
                    c[i, j] = c[i-1, j]
                    b[i, j] = 1
                else:
                    c[i, j] = c[i, j-1]
                    b[i, j] = 3
    return c, b

def get_max_subseq(texta, textb):
    """
    获取最长公共序列
    如：
        x = '荣丰控股集团股份有限公司'
        y = '荣丰（天津）医疗器械有限公司'
    return:
        荣丰有限公司
    :param x:
    :param y:
    :return:
    """
    c, b = LCS(texta,textb)
    i = len(texta)
    j = len(textb)
    lcs = ''
    while i > 0 and j > 0:
        if b[i][j] == 2:
            lcs = texta[i-1]+lcs
            i -= 1
            j -= 1
        if b[i][j] == 1:
            i -= 1
        if b[i][j] == 3:
            j -= 1
        if b[i][j] == 0:
            break
    return lcs

def get_max_comstr(str1,str2):
    """
    最长公共字串
    动态规划
    时间复杂度O(n^2)
        例： str1 = '北京人爱喝豆汁？'
            str2 = '老北京喜欢喝豆汁？'
            sub = max_sub_str(str1,str2)
            结果：sub = (4, ['喝豆汁？'])
    :param str1:
    :param str2:
    :return: 最大长度，最大长度子串列表
    """
    res = [[0]*len(str2) for _ in range(len(str1))]
    max_len=0
    max_str=[]
    for i in range(len(str1)):
        for j in range(len(str2)):
            if str1[i]==str2[j]:
                if i>0 and j>0:
                    res[i][j]=1+res[i-1][j-1]
                else:
                    res[i][j]=1

            # 记录最大值
            if res[i][j]>max_len:
                max_len=res[i][j]
                max_str=[str1[i-max_len+1:i+1]]
            elif res[i][j]==max_len:
                max_str.append(str1[i-max_len:i])

    return max_len,max_str

def is_chinese(string):
    """
    检查整个字符串是否包含中文，若包含则返回True
    :param string: 需要检查的字符串
    :return: bool
    """
    for ch in string:
        if u'\u4e00' <= ch <= u'\u9fff':
            return True

    return False

def change_strQ2B(ustring):
    """
    字符串全角转半角
    """
    rstring = ""
    for uchar in ustring:
        inside_code = ord(uchar)
        if inside_code == 12288:  #全角空格直接转换
            inside_code = 32
        elif (inside_code >= 65281 and inside_code <= 65374):  #全角字符（除空格）根据关系转化
            inside_code -= 65248

        rstring += chr(inside_code)
    return rstring

def read_json(filename):
    """
    读取json文件
    :param filename: json文件路径
    :return: json格式内容
    """
    try:
        with open(filename, 'r', encoding='utf8') as f:
            json_data = json.load(f)
        return json_data
    except:
        json_data = []
        with open(filename, encoding='utf-8') as f:
            for line in tqdm(f):
                json_data.append(json.loads(line))
        return json_data

class MyEncoder(json.JSONEncoder):
    def default(self, obj):
        if isinstance(obj, np.integer):
            return int(obj)
        elif isinstance(obj, np.floating):
            return float(obj)
        elif isinstance(obj, np.ndarray):
            return obj.tolist()
        elif isinstance(obj, bytes):
            return str(obj, encoding='utf-8')
        elif isinstance(obj, datetime.datetime):
            return obj.strftime('%Y-%m-%d %H:%M:%S')
        elif isinstance(obj, datetime.date):
            return obj.strftime('%Y-%m-%d')
        else:
            return super(MyEncoder, self).default(obj)

def write_json(filename, data, isIndent=True, isLine=False):
    """
    将字典或者list数据保存为json
    :param filename: 保存的文件名
    :param data: 要保存的数据
    :param isIndent: 是否按照漂亮的格式保存
    :param isLine: 是否按行保存
    :return:
    """
    if not isLine and (isinstance(data, dict) or isinstance(data, list)):
        if isIndent:
            json_str = json.dumps(data, ensure_ascii=False, indent=4, cls=MyEncoder)
            with open(filename, 'w', encoding='utf-8') as json_file:
                json_file.write(json_str)
        else:
            json_str = json.dumps(data, ensure_ascii=False, cls=MyEncoder)
            with open(filename, 'w', encoding='utf-8') as json_file:
                json_file.write(json_str)
    else:
        with codecs.open(filename, 'w', 'utf-8') as f:
            for formatted_instance in data:
                json_str = json.dumps(formatted_instance, ensure_ascii=False, cls=MyEncoder)
                f.write(json_str)
                f.write('\n')
        f.close()


def read_csv(file_path, isdict=False):
    """
    读取csv数据
    返回数据格式如：
        isdict=False（默认）：
            list，且每个元素也是list，每个元素表示每行数据
            [[line1_cell1, line1_cell2], [line2_cell1, line2_cell2], ...]
                例如：
                [
                    ['class','name','sex','height','year'],
                    [1,'xiaoming','male',168,23],
                    [2,'xiaohong','female',162,22],
                    [3,'xiaozhang','female',163,21],
                    [4,'xiaoli','male',158,21],
                    ...
                ]
        isdict=True：
            list，每个元素是dict，每个元素表示每行数据
            [{key1: line1_cell1, key2, line1_cell2, ...}, {key1: line2_cell1, key2, line2_cell2, ...}, ...]
                例如：
                [
                    {'class': '1', 'name': 'xiaoming', 'sex': 'male', 'height': '168', year: '23'},
                    {'class': '2', 'name': 'xiaohong', 'sex': 'female', 'height': '162', year: '22'},
                    {'class': '3', 'name': 'xiaozhang', 'sex': 'female', 'height': '163', year: '21'},
                    {'class': '4', 'name': 'xiaoli', 'sex': 'male', 'height': '158', year: '21'},
                    ...
                ]
    :param file_path:csv文件路径
    :param isdict: 返回数据格式，默认False， 返回的每行作为一个list， 如果设为True，则每行作为一个dict
    :return: list， 根据isdict觉得每个元素的格式
    """
    res = []
    with open(file_path, encoding='utf-8') as f:
        reader = csv.reader(f)
        try:
            for l in reader:
                if len(l) == 0:
                    continue
                res.append(l)
        except Exception as e:
            print("\033[1;31m 警告：读取csv时发生错误，已经读取到：{} 条数据， 错误内容： {}\033[0m".format(len(res), e))

    if isdict:
        if len(res) < 1:
            return []

        keys = []
        keys_dict = collections.defaultdict(int)
        for k in res[0]:
            k = str(k)
            keys_dict[k] += 1
            if k not in keys:
                keys.append(k)
            else:
                k_new = k + '_' + str(keys_dict[k])
                keys.append(k_new)

        res_dict = []
        for d in res[1:]:
            linei = {}
            for di, ki in zip(d, keys):
                linei[ki] = di
            res_dict.append(linei)

        return res_dict
    return res


def write_csv(filename, data):
    """
    将数据写到csv中
    支持两种数据格式：
        （1） list，且每个元素也是list，每个元素表示每行数据
            [[line1_cell1, line1_cell2], [line2_cell1, line2_cell2], ...]
        例如：
            [
                ['class','name','sex','height','year'],
                [1,'xiaoming','male',168,23],
                [2,'xiaohong','female',162,22],
                [3,'xiaozhang','female',163,21],
                [4,'xiaoli','male',158,21],
                ...
            ]

        （2） list，每个元素是dict，每个元素表示每行数据
            [{key1: line1_cell1, key2, line1_cell2, ...}, {key1: line2_cell1, key2, line2_cell2, ...}, ...]
        例如：
            [
                {'class': '1', 'name': 'xiaoming', 'sex': 'male', 'height': '168', 'year': '23'},
                {'class': '2', 'name': 'xiaohong', 'sex': 'female', 'height': '162', 'year': '22'},
                {'class': '3', 'name': 'xiaozhang', 'sex': 'female', 'height': '163', 'year': '21'},
                {'class': '4', 'name': 'xiaoli', 'sex': 'male', 'height': '158', 'year': '21'},
                ...
            ]
    :param filename: 需要写入的csv文件路径
    :param data: 需要写入的数据
    """
    isdict = False
    if len(data) > 0:
        if type(data[0]) == dict:
            isdict = True
    f = open(filename, 'w', encoding='utf-8', newline='')
    write_data = []
    keys = []
    if isdict:
        for d in data:
            keysi = list(d.keys())
            for ki in keysi:
                if ki not in keys:
                    keys.append(ki)

        write_data.append(keys)
        for d in data:
            di = []
            for k in keys:
                if k in d:
                    di.append(d[k])
                else:
                    di.append('')
            write_data.append(di)
    else:
        write_data = data
    writer = csv.writer(f)
    for i in write_data:
        writer.writerow(i)
    f.close()


def read_excel(filename, sheet_name='', isdict=False):
    """
    读取excel数据， 默认读取第一个sheet，可以通过sheet_name决定读取的sheet
    返回数据格式如：
        isdict=False（默认）：
            list，且每个元素也是list，每个元素表示每行数据
            [[line1_cell1, line1_cell2], [line2_cell1, line2_cell2], ...]
                例如：
                [
                    ['class','name','sex','height','year'],
                    [1,'xiaoming','male',168,23],
                    [2,'xiaohong','female',162,22],
                    [3,'xiaozhang','female',163,21],
                    [4,'xiaoli','male',158,21],
                    ...
                ]
        isdict=True：
            list，每个元素是dict，每个元素表示每行数据
            [{key1: line1_cell1, key2, line1_cell2, ...}, {key1: line2_cell1, key2, line2_cell2, ...}, ...]
                例如：
                [
                    {'class': '1', 'name': 'xiaoming', 'sex': 'male', 'height': '168', year: '23'},
                    {'class': '2', 'name': 'xiaohong', 'sex': 'female', 'height': '162', year: '22'},
                    {'class': '3', 'name': 'xiaozhang', 'sex': 'female', 'height': '163', year: '21'},
                    {'class': '4', 'name': 'xiaoli', 'sex': 'male', 'height': '158', year: '21'},
                    ...
                ]
    :param filename: excel文件路径
    :param sheet_name: 需要读取的excel中sheet的名称， 默认读取第一个sheet
    :param isdict: 返回数据格式，默认False， 返回的每行作为一个list， 如果设为True，则每行作为一个dict
    :return: list， 根据isdict觉得每个元素的格式
    """
    res = []

    data = xlrd.open_workbook(filename)

    if sheet_name != "":
        table = data.sheet_by_name(sheet_name)
    else:
        table = data.sheets()[0]

    rowNum = table.nrows
    colNum = table.ncols

    for i in range(rowNum):
        row_data = []
        for j in range(colNum):
            cell_ij = table.cell(i, j)

            value = table.cell(i, j).value
            if cell_ij.ctype == 4:
                if value == 1:
                    value = True
                else:
                    value = False

            if value == 'null':
                value = ''
            row_data.append(value)
        res.append(row_data)

    if isdict:
        if len(res) < 1:
            return None

        keys = []
        keys_dict = collections.defaultdict(int)
        for k in res[0]:
            k = str(k)
            keys_dict[k] += 1
            if k not in keys:
                keys.append(k)
            else:
                k_new = k + '_' + str(keys_dict[k])
                keys.append(k_new)

        res_dict = []
        for d in res[1:]:
            linei = {}
            for di, ki in zip(d, keys):
                linei[ki] = di
            res_dict.append(linei)

        return res_dict

    return res


def write_excel(filename, data, sheet_name='Sheet1'):
    """
    将数据写到excel中， 默认sheet_name='Sheet1'， 可以自行设置
    支持两种数据格式：
        （1） list，且每个元素也是list，每个元素表示每行数据
            [[line1_cell1, line1_cell2], [line2_cell1, line2_cell2], ...]
        例如：
            [
                ['class','name','sex','height','year'],
                [1,'xiaoming','male',168,23],
                [2,'xiaohong','female',162,22],
                [3,'xiaozhang','female',163,21],
                [4,'xiaoli','male',158,21],
                ...
            ]

        （2） list，每个元素是dict，每个元素表示每行数据
            [{key1: line1_cell1, key2, line1_cell2, ...}, {key1: line2_cell1, key2, line2_cell2, ...}, ...]
        例如：
            [
                {'class': '1', 'name': 'xiaoming', 'sex': 'male', 'height': '168', year: '23'},
                {'class': '2', 'name': 'xiaohong', 'sex': 'female', 'height': '162', year: '22'},
                {'class': '3', 'name': 'xiaozhang', 'sex': 'female', 'height': '163', year: '21'},
                {'class': '4', 'name': 'xiaoli', 'sex': 'male', 'height': '158', year: '21'},
                ...
            ]
    :param filename: 需要写入的excel文件路径， 如'a.xlsx'
    :param data: 需要写入的数据
    :param sheet_name: sheet的名称， 默认为：Sheet1
    """
    isdict = False
    if len(data) > 0:
        if type(data[0]) == dict:
            isdict = True
    workbook = xlsxwriter.Workbook(filename)  # 创建一个excel文件
    worksheet = workbook.add_worksheet(sheet_name)  # 在文件中创建一个名为TEST的sheet,不加名字默认为sheet1
    write_data = []
    keys = []
    if isdict:
        for d in data:
            keysi = list(d.keys())
            for ki in keysi:
                if ki not in keys:
                    keys.append(ki)

        write_data.append(keys)
        for d in data:
            di = []
            for k in keys:
                if k in d:
                    di.append(d[k])
                else:
                    di.append('')
            write_data.append(di)
    else:
        write_data = data

    if len(write_data) > 0:
        for i in range(len(write_data)):
            for j in range(len(write_data[i])):
                worksheet.write(i, j, write_data[i][j])

    workbook.close()


def read_txt(fileName):
    """
        读取txt文件， 返回list， 每行内容为元素
    """
    res = []
    with open(fileName, "r" ,encoding='utf8') as f:
        data = f.readlines()
        for d in tqdm(data):
            res.append(d.strip('\n'))

    return res

def write_txt(fileName, text_list, model='normal'):
    """
        将文本list写入到txt文件中，每个元素一行内容
    :param fileName: 写入的文件名
    :param text_list: 字符串组成的list
    :param model: 默认normal，新建txt并写入，如果其他，则为新增
    """
    if model == 'normal':
        f = open(fileName, 'w', encoding='utf8')
    else:
        f = open(fileName, 'a', encoding='utf8')
    for i in tqdm(text_list):
        f.write(str(i))
        f.write('\n')
    f.close()


def iter_block_items(parent):
    """
        用于读取word文档
    """
    if isinstance(parent, Document):
        parent_elm = parent.element.body
    elif isinstance(parent, _Cell):
        parent_elm = parent._tc
    else:
        raise ValueError("something's not right")

    for child in parent_elm.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, parent)
        elif isinstance(child, CT_Tbl):
            yield Table(child, parent)

def read_table(table):
    return [[cell.text for cell in row.cells] for row in table.rows]

def read_word(word_path):
    """
        读取docx文档，读取其中的文本和表格，结果组成一个list
        如果是doc结尾，则需要先转换为docx格式，才能读取
        无法获取其他非文本内容（图片、音频等）
    :param word_path:word文档路径
    :return:文本和表格组成的list
    """
    res = []
    doc = docx.Document(word_path)
    for block in iter_block_items(doc):
        if isinstance(block, Paragraph):
            res.append(block.text)
        elif isinstance(block, Table):
            res.append(read_table(block))
    return res


def compute_vecsimilar_one_2_one(veca, vecb):
    """
    计算一个向量和另一个向量之间的相似度
    :param veca: 第一个向量
    :param vecb: 第二个向量
    :return: 两个向量的相似度
    """
    veca = np.array(veca)
    vecb = np.array(vecb)
    veca = veca / (veca ** 2).sum() ** 0.5
    vecb = vecb / (vecb ** 2).sum() ** 0.5
    sim = (veca * vecb).sum()
    return sim


def compute_vecsimilar__one_2_many(veca, vecs):
    """
    计算一条数据和多个向量之间的相似度
    分别返回：
        最大相似度下标: maxarg
        最大相似度对应相似度：maxsim
    :param veca: vec
    :param vecs:[vec1, vec2, ...]
    :return:maxarg, maxsim
    """

    veca = np.array(veca)
    vecs = np.array(vecs)
    veca = veca / (veca ** 2).sum() ** 0.5
    vecs = vecs / (vecs ** 2).sum(axis=1, keepdims=True) ** 0.5

    sims = np.dot(veca, vecs.T)
    maxarg = sims.argmax()
    maxsim = max(sims)
    return maxarg, maxsim


def compute_vecsimilar__many_2_many(vecsa, vecsb):
    """
    计算多个向量和多个向量之间的相似度
    分别返回：
        maxarg： vecsa中的每条和vecsb中相似度最高的下标
        maxsim： vecsa中的每条和vecsb中相似度最高的相似度
    :param vecsa:[vec1, vec2, ...]
    :param vecsb:[vec1, vec2, ...]
    :return:maxargs, maxsims
    """
    veca = np.array(vecsa)
    vecs = np.array(vecsb)
    veca = veca / (veca ** 2).sum(axis=1, keepdims=True) ** 0.5
    vecs = vecs / (vecs ** 2).sum(axis=1, keepdims=True) ** 0.5

    sims = np.dot(veca, vecs.T)

    maxargs = sims.argmax(1)
    maxsims = sims.max(1)
    return maxargs, maxsims




